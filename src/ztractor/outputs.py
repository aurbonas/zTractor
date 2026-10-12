"""Gravacao dos resultados em JSON, CSV ou Excel.

No Excel, a aba `results` tem uma linha por tarefa e cada tarefa que devolveu
uma lista (jobs, datasets, membros, arquivos USS) ganha sua propria aba.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import List

from .engine import TaskResult


def _items(data):
    """Extrai a lista de registros de uma resposta do Zowe, se houver."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ("items", "apiResponse"):
            value = data.get(key)
            if isinstance(value, list):
                return value
            if isinstance(value, dict) and isinstance(value.get("items"), list):
                return value["items"]
    return None


def write_json(results: List[TaskResult], path: Path) -> None:
    payload = [dict(r.as_row(), data=r.data) for r in results]
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")


def write_csv(results: List[TaskResult], path: Path) -> None:
    rows = [r.as_row() for r in results]
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()) if rows else ["n"], delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def _cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def write_xlsx(results: List[TaskResult], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    rows = [r.as_row() for r in results]
    header = list(rows[0].keys()) if rows else ["n"]
    ws.append(header)
    for row in rows:
        ws.append([_cell(row[h]) for h in header])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in results:
        items = _items(r.data)
        if not items or not all(isinstance(i, dict) for i in items):
            if isinstance(r.data, str) and r.data.strip():
                detail = wb.create_sheet(f"{r.index}_{r.action}"[:31])
                for line in r.data.splitlines():
                    detail.append([line])
            continue
        detail = wb.create_sheet(f"{r.index}_{r.action}"[:31])
        cols: List[str] = []
        for item in items:
            for k in item.keys():
                if k not in cols:
                    cols.append(k)
        detail.append(cols)
        for item in items:
            detail.append([_cell(item.get(c)) for c in cols])
        for cell in detail[1]:
            cell.font = Font(bold=True)
    wb.save(path)


def write(results: List[TaskResult], path: str) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        write_xlsx(results, p)
    elif ext == ".csv":
        write_csv(results, p)
    else:
        write_json(results, p)
