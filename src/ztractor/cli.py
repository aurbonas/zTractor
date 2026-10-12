"""Linha de comando do zTractor.

    ztractor actions                         lista as acoes
    ztractor plan  ENTRADA                   mostra os comandos sem executar
    ztractor run   ENTRADA [-o saida.xlsx]   executa
    ztractor do    ACAO [chave=valor ...]    executa uma acao avulsa
    ztractor template saida.xlsx             gera planilha modelo

ENTRADA: .xlsx, .csv, .json, .yaml, .txt ou URL http(s).
"""

from __future__ import annotations

import argparse
import json
import sys
from typing import Dict, List

from . import __version__
from .actions import ACTIONS, DESTRUTIVA
from .engine import FALHOU, INVALIDO, PULADO, run_plan
from .inputs import Plan, load
from .inputs.rest import post_results
from .inputs.text import parse_text
from .outputs import write


def _headers(values: List[str]) -> Dict[str, str]:
    out = {}
    for h in values or []:
        if ":" not in h:
            raise SystemExit(f"cabecalho invalido: {h} (use 'Nome: valor')")
        k, v = h.split(":", 1)
        out[k.strip()] = v.strip()
    return out


def _vars(values: List[str]) -> Dict[str, str]:
    out = {}
    for item in values or []:
        if "=" not in item:
            raise SystemExit(f"--var invalido: {item} (use NOME=valor)")
        k, v = item.split("=", 1)
        out[k] = v
    return out


def cmd_actions(args) -> int:
    rows = sorted(ACTIONS.values(), key=lambda a: (a.family, a.name))
    if args.json:
        print(json.dumps([{"action": a.name, "family": a.family, "risk": a.risk, "required": a.required,
                           "summary": a.summary} for a in rows], indent=2, ensure_ascii=False))
        return 0
    for a in rows:
        req = ", ".join(a.required) or "-"
        print(f"{a.name:<20} {a.risk:<10} campos: {req:<22} {a.summary}")
    return 0


def _finish(results, args) -> int:
    rows = [dict(r.as_row(), data=r.data) for r in results]
    if args.output:
        write(results, args.output)
        print(f"resultados gravados em {args.output}", file=sys.stderr)
    if getattr(args, "post_results", None):
        post_results(args.post_results, rows, _headers(args.header))
        print(f"resultados enviados para {args.post_results}", file=sys.stderr)
    if not args.output or args.json:
        print(json.dumps(rows, indent=2, ensure_ascii=False, default=str))
    bad = [r for r in results if r.status in (FALHOU, INVALIDO)]
    skipped = [r for r in results if r.status == PULADO]
    print(f"{len(results)} tarefas: {len(results) - len(bad) - len(skipped)} ok/planejadas, "
          f"{len(bad)} com erro, {len(skipped)} puladas", file=sys.stderr)
    return 1 if bad else (2 if skipped else 0)


def cmd_run(args, dry_run: bool) -> int:
    plan = load(args.input, sheet=args.sheet, headers=_headers(args.header))
    plan.vars.update(_vars(args.var))
    if args.only:
        wanted = {s.strip().lower() for s in args.only.split(",")}
        plan.tasks = [t for t in plan.tasks if str(t.get("action", "")).lower() in wanted
                      or str(t.get("action", "")).split(".")[0].lower() in wanted]
    results = run_plan(plan, dry_run=dry_run, yes=args.yes, stop_on_error=args.stop_on_error,
                       profile=args.profile)
    return _finish(results, args)


def cmd_do(args) -> int:
    line = " ".join([args.action] + [f'"{p}"' if " " in p else p for p in args.params])
    variables, tasks = parse_text(line)
    plan = Plan(variables, tasks, "linha de comando")
    plan.vars.update(_vars(args.var))
    results = run_plan(plan, dry_run=args.dry_run, yes=args.yes, profile=args.profile)
    return _finish(results, args)


def cmd_template(args) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "tasks"
    header = ["action", "target", "enabled", "owner", "prefix", "to", "file", "dir", "content",
              "command", "type", "recfm", "lrecl", "size", "wait", "replace", "profile", "obs"]
    ws.append(header)
    examples = [
        ["setup.check", None, "sim"],
        ["jobs.list", None, "sim", "${USER_ZOS}", "DAILY*"],
        ["ds.list", "${HLQ}.**", "sim"],
        ["ds.members", "${HLQ}.JCL", "sim"],
        ["ds.create", "${HLQ}.ZTRACTOR.DATA", "nao", None, None, None, None, None, None, None,
         "seq", "FB", 80, "1CYL"],
        ["ds.write", "${HLQ}.ZTRACTOR.DATA", "nao", None, None, None, None, None, "LINHA 1\\nLINHA 2"],
        ["jobs.submit", "${HLQ}.JCL(IEFBR14)", "nao", None, None, None, None, None, None, None,
         None, None, None, None, "sim"],
        ["console.issue", None, "sim", None, None, None, None, None, None, "D T"],
        ["tso.issue", None, "nao", None, None, None, None, None, None, "LISTCAT LEVEL(${HLQ})"],
    ]
    for row in examples:
        ws.append(row + [None] * (len(header) - len(row)))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wv = wb.create_sheet("vars")
    wv.append(["name", "value"])
    wv.append(["HLQ", "Z00000"])
    wv.append(["USER_ZOS", "Z00000"])
    for cell in wv[1]:
        cell.font = Font(bold=True)
    wa = wb.create_sheet("acoes")
    wa.append(["action", "risco", "campos obrigatorios", "o que faz"])
    for a in sorted(ACTIONS.values(), key=lambda a: (a.family, a.name)):
        wa.append([a.name, a.risk, ", ".join(a.required), a.summary])
    for cell in wa[1]:
        cell.font = Font(bold=True)
    wb.save(args.path)
    print(f"modelo gravado em {args.path}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="ztractor", description="z/OS com Zowe CLI a partir de Excel, REST e texto")
    p.add_argument("--version", action="version", version=f"ztractor {__version__}")
    sub = p.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("actions", help="lista as acoes disponiveis")
    a.add_argument("--json", action="store_true")

    for name, helptext in (("plan", "mostra os comandos sem executar"), ("run", "executa as tarefas")):
        r = sub.add_parser(name, help=helptext)
        r.add_argument("input", help="arquivo .xlsx/.csv/.json/.yaml/.txt ou URL")
        r.add_argument("--sheet", help="aba da planilha com as tarefas (padrao: tasks)")
        r.add_argument("--header", action="append", default=[], help="cabecalho HTTP 'Nome: valor' (aceita ${VAR})")
        r.add_argument("--var", action="append", default=[], help="variavel NOME=valor")
        r.add_argument("--only", help="so estas acoes ou familias, separadas por virgula")
        r.add_argument("--profile", help="perfil zosmf padrao")
        r.add_argument("-o", "--output", help="grava resultados (.xlsx, .csv ou .json)")
        r.add_argument("--json", action="store_true", help="imprime o JSON mesmo com --output")
        r.add_argument("--post-results", help="envia os resultados por POST para esta URL")
        r.add_argument("--yes", action="store_true", help=f"confirma acoes de risco {DESTRUTIVA}")
        r.add_argument("--stop-on-error", action="store_true")

    d = sub.add_parser("do", help="executa uma acao avulsa: ztractor do ds.list HLQ.**")
    d.add_argument("action")
    d.add_argument("params", nargs="*", help="alvo e chave=valor")
    d.add_argument("--var", action="append", default=[])
    d.add_argument("--profile")
    d.add_argument("--dry-run", action="store_true")
    d.add_argument("--yes", action="store_true")
    d.add_argument("-o", "--output")
    d.add_argument("--json", action="store_true")

    t = sub.add_parser("template", help="gera uma planilha modelo")
    t.add_argument("path", nargs="?", default="ztractor-tarefas.xlsx")
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.cmd == "actions":
            return cmd_actions(args)
        if args.cmd == "plan":
            return cmd_run(args, dry_run=True)
        if args.cmd == "run":
            return cmd_run(args, dry_run=False)
        if args.cmd == "do":
            return cmd_do(args)
        if args.cmd == "template":
            return cmd_template(args)
    except (FileNotFoundError, ValueError, KeyError) as exc:
        print(f"erro: {str(exc).strip(chr(39))}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
