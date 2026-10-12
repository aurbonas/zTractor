"""Planilha Excel de entrada.

Aba de tarefas: a aba `tasks` (ou a indicada em --sheet, ou a primeira aba).
A primeira linha e o cabecalho; cada linha seguinte e uma tarefa.
Aba `vars` opcional: colunas `name` e `value`.
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


def _rows(ws) -> List[dict]:
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    keys = [str(h).strip() if h is not None else None for h in header]
    out = []
    for row in rows:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        out.append({k: v for k, v in zip(keys, row) if k})
    return out


def read_excel(path: Path, sheet: Optional[str] = None) -> Tuple[Dict[str, str], List[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    names = {n.lower(): n for n in wb.sheetnames}
    if sheet:
        if sheet.lower() not in names:
            raise ValueError(f"aba '{sheet}' nao existe em {path.name}: {', '.join(wb.sheetnames)}")
        ws = wb[names[sheet.lower()]]
    elif "tasks" in names:
        ws = wb[names["tasks"]]
    else:
        ws = wb[wb.sheetnames[0]]
    tasks = _rows(ws)
    variables: Dict[str, str] = {}
    if "vars" in names and ws.title.lower() != "vars":
        for row in _rows(wb[names["vars"]]):
            low = {str(k).lower(): v for k, v in row.items()}
            name = low.get("name") or low.get("nome")
            value = low.get("value", low.get("valor"))
            if name:
                variables[str(name).strip()] = "" if value is None else str(value)
    wb.close()
    return variables, tasks
