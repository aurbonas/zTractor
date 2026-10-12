from openpyxl import load_workbook

from ztractor.cli import main
from ztractor.engine import FALHOU, INVALIDO, OK, PLANEJADO, PULADO, run_plan
from ztractor.inputs import Plan


def plan(*tasks, **vars_):
    return Plan(dict(vars_), list(tasks))


def test_dry_run_calls_nothing(fake_zowe):
    res = run_plan(plan({"action": "ds.list", "target": "${HLQ}.**"}, HLQ="Z1"), dry_run=True)
    assert res[0].status == PLANEJADO
    assert '"Z1.**"' in res[0].command
    assert fake_zowe() == []


def test_run_parses_rfj_after_warning(fake_zowe):
    res = run_plan(plan({"action": "jobs.list", "owner": "Z1"}))
    assert res[0].status == OK
    assert res[0].data[0]["jobid"] == "JOB00001"
    assert fake_zowe()[0]["args"] == ["zos-jobs", "list", "jobs", "--owner", "Z1", "--rfj"]


def test_destructive_needs_yes(fake_zowe):
    t = {"action": "ds.delete", "target": "Z1.OLD"}
    res = run_plan(plan(t), confirm=lambda r: False)
    assert res[0].status == PULADO and fake_zowe() == []
    res = run_plan(plan(t), yes=True)
    assert res[0].status == OK
    assert fake_zowe()[0]["args"][-2:] == ["--for-sure", "--rfj"]


def test_failure_and_invalid(fake_zowe):
    res = run_plan(plan({"action": "ds.view", "target": "FAIL.X"}, {"action": "ds.view"}, {"action": "jobs.list"}),
                   stop_on_error=False)
    assert [r.status for r in res] == [FALHOU, INVALIDO, OK]
    assert res[0].message == "falha simulada"


def test_write_sends_stdin(fake_zowe):
    run_plan(plan({"action": "ds.write", "target": "Z1.D", "content": "A\\nB"}), yes=True)
    assert fake_zowe()[0]["stdin"] == "A\nB\n"


def test_submit_renders_local_jcl(fake_zowe, tmp_path):
    jcl = tmp_path / "a.jcl"
    jcl.write_text("//J JOB\n//S EXEC PGM=${PGM}\n")
    run_plan(plan({"action": "jobs.submit", "target": str(jcl)}, PGM="IEFBR14"))
    call = fake_zowe()[0]["args"]
    assert call[:3] == ["zos-jobs", "submit", "local-file"]
    assert call[3] != str(jcl)  # arquivo temporario ja substituido


def test_cli_run_to_excel(fake_zowe, tmp_path):
    src = tmp_path / "t.txt"
    src.write_text("jobs.list owner=Z1\nds.list Z1.**\nds.view Z1.JCL(A)\n")
    out = tmp_path / "r.xlsx"
    assert main(["run", str(src), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["results", "1_jobs.list", "2_ds.list", "3_ds.view"]
    assert wb["2_ds.list"]["A2"].value == "Z1.JCL"
    assert wb["results"]["E2"].value == "ok"


def test_cli_skips_destructive_without_yes(fake_zowe, tmp_path):
    src = tmp_path / "t.txt"
    src.write_text("jobs.cancel JOB00001\n")
    assert main(["run", str(src), "-o", str(tmp_path / "r.json")]) == 2
